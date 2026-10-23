# -*- coding: UTF-8 -*-
# @Author     zyb
# @FileName   init_voice.PY
# @Time       2025-05-30 22:56

import time
import shutil
import datetime
from pydub import AudioSegment
from openpyxl import load_workbook
from typing import Dict, Optional, Any
from conf.basic import *
from utils.recognition_by_kedaxunfei import make_tts_kedaxunfei

logger = get_logger("face_v")



def add_spe(user_id: str, speech1: str, speech2: str) -> Dict[str, Any]:
    try:
        # 定义配置文件路径常量
        CONF_DIR = os.path.join(os.path.dirname(__file__), "conf")
        JSON_PATH = os.path.join(CONF_DIR, "speech_json.json")

        # 确保配置目录存在
        os.makedirs(CONF_DIR, exist_ok=True)

        # 加载现有数据
        data: Dict[str, Dict[str, Optional[str]]] = {}
        if os.path.exists(JSON_PATH):
            with open(JSON_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)

        # 查找现有用户ID对应的语音ID（优化查找效率）
        existing_id = next(
            (speech_id for speech_id, content in data.items()
             if content.get("user_id") == user_id),
            None
        )

        if existing_id:
            # 更新现有用户的语音数据
            data[existing_id].update({
                "speech1": speech1,
                "speech2": speech2
            })
            logger.info(f"成功更新用户[{user_id}]的语音数据: {existing_id}")
        else:
            # 为新用户创建语音数据
            speech_id = f"{int(time.time())}"
            data[speech_id] = {
                "user_id": user_id,
                "speech1": speech1,
                "speech2": speech2
            }
            logger.info(f"成功为用户[{user_id}]添加语音数据: {speech_id}")

        # 保存更新后的数据
        with open(JSON_PATH, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=4)

        return {"code": 200, "message": "语音数据操作成功"}

    except Exception as e:
        error_msg = f"语音数据操作失败: {str(e)}"
        logger.error(error_msg)
        return {"code": 500, "message": error_msg}

def add_failed(failed: str):
    try:
        conf_dir = os.path.join(os.path.dirname(__file__), "conf")
        os.makedirs(conf_dir, exist_ok=True)
        json_path = os.path.join(conf_dir, "failed_json.json")
        data = {
            f"{int(time.time())}": {
                "failed": failed
            }
        }
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=4)
        logger.info(f"成功更新失败语音数据")
        return {"code": 200, "message": "失败语音数据更新成功"}
    except Exception as e:
        logger.error(f"更新失败语音数据失败: {str(e)}")
        return {"code": 500, "message": f"更新失败语音数据失败: {str(e)}"}

def remake_path():
    if os.path.isdir(GlobalConf.voice_path):
        shutil.rmtree(GlobalConf.voice_path)
    os.makedirs(GlobalConf.voice_path)

def all_voice(speech1, speech2, filename):
    first_voice = speech1
    second_voice = speech2
    excel_file = os.path.join("name_file", filename)
    workbook = load_workbook(excel_file)
    sheet = workbook.active
    logger.info(f"当前统计到的人数:{sheet.max_row}")
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row):
        name = row[0].value
        if all('\u4e00' <= char <= '\u9fa5' for char in name):
            run_first_message = f"{name}{first_voice}"
            run_second_message = f"{name}{second_voice}"
        else:
            run_first_message = f"{first_voice}"
            run_second_message = f"{second_voice}"
        voice_file = os.path.join(GlobalConf.voice_path, f"{name}_1.mp3")
        logger.info(f"生成语音:{run_first_message},文件:{voice_file}")
        while 1:
            flag = make_tts_kedaxunfei(logger, GlobalConf.app_id, GlobalConf.api_key, GlobalConf.vcn, run_first_message,
                                       voice_file)
            if flag:
                break
        if not os.path.isfile(voice_file):
            logger.error(f"{name}没有语音:{voice_file}")
        audio = AudioSegment.from_mp3(voice_file)
        wav_path = os.path.join(GlobalConf.voice_path, f"{name}_1.wav")
        audio.export(wav_path, format="wav")

        voice_file = os.path.join(GlobalConf.voice_path, f"{name}_2.mp3")
        logger.info(f"生成语音:{run_second_message},文件:{voice_file}")
        while 1:
            flag = make_tts_kedaxunfei(logger, GlobalConf.app_id, GlobalConf.api_key, GlobalConf.vcn,
                                       run_second_message,
                                       voice_file)
            if flag:
                break
        if not os.path.isfile(voice_file):
            logger.error(f"{name}没有语音:{voice_file}")
        audio = AudioSegment.from_mp3(voice_file)
        wav_path = os.path.join(GlobalConf.voice_path, f"{name}_2.wav")
        audio.export(wav_path, format="wav")


def add_one(user_id, name, speech1, speech2):
    voice_dir = os.path.join(os.path.dirname(__file__), "voice")
    os.makedirs(voice_dir, exist_ok=True)
    greetings = [
        speech1,
        speech2
    ]
    for idx, message in enumerate(greetings, start=1):
        voice_file = os.path.join(voice_dir, f"{name}_{idx}.mp3")
        # if all('\u4e00' <= char <= '\u9fff' for char in name):
        if all('\u4e00' <= char <= '\u9fa5' for char in name):
            full_message = f"{name}{message}"
        else:
            full_message = f"{message}"
        logger.info(f"正在为 {name} 生成第{idx}段语音: {voice_file}")
        while True:
            success = make_tts_kedaxunfei(
                logger,
                GlobalConf.app_id,
                GlobalConf.api_key,
                GlobalConf.vcn,
                full_message,
                voice_file
            )
            if success:
                break
        if os.path.exists(voice_file):
            audio = AudioSegment.from_mp3(voice_file)
            wav_path = os.path.join(voice_dir, f"{name}_{idx}.wav")
            audio.export(wav_path, format="wav")
            add_spe(user_id, speech1, speech2)
            logger.info(f"成功生成WAV文件: {wav_path}")
        else:
            logger.error(f"生成MP3文件失败: {voice_file}")

def add_fix():
    run_message = "请正视摄像头"
    voice_file = os.path.join(GlobalConf.voice_path, f"show.mp3")
    while 1:
        flag = make_tts_kedaxunfei(logger, GlobalConf.app_id, GlobalConf.api_key, GlobalConf.vcn,
                                   run_message,
                                   voice_file)
        if flag:
            break

def add_fail(text):
    voice_file = os.path.join(GlobalConf.voice_path, f"fail.mp3")
    while 1:
        flag = make_tts_kedaxunfei(logger, GlobalConf.app_id, GlobalConf.api_key, GlobalConf.vcn,
                                   text,
                                   voice_file)
        if flag:
            break
    if os.path.exists(voice_file):
        audio = AudioSegment.from_mp3(voice_file)
        wav_path = os.path.join(GlobalConf.voice_path, "fail.wav")
        audio.export(wav_path, format="wav")
        os.remove(voice_file)
        add_failed(text)
        logger.info(f"已转换并删除MP3文件，生成WAV文件: {wav_path}")

def add_voice(anystr):
    run_message = anystr
    voice_name = datetime.datetime.now().strftime('%Y%m%d%H%M%S')
    parent_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    mp3_dir = os.path.join(parent_dir, 'mp3')
    voice_file = os.path.join(mp3_dir, f"{voice_name}.mp3")
    while 1:
        flag = make_tts_kedaxunfei(logger, GlobalConf.app_id, GlobalConf.api_key, GlobalConf.vcn, run_message, voice_file)
        if flag:
            return voice_file

