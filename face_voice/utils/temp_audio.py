# -*- coding: UTF-8 -*-
# @Author     zyb
# @FileName   temp_audio.PY
# @Time       2025-07-01 09:55
import os
import random
from openpyxl import load_workbook
from conf.basic import *
from pydub import AudioSegment
from utils.recognition_by_kedaxunfei import make_tts_kedaxunfei


logger = get_logger("temp_audio")

def down_voice(name, massage):
    temp_mp3_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "conf", "temp_mp3")
    os.makedirs(temp_mp3_dir, exist_ok=True)
    voice_file = os.path.join(temp_mp3_dir, f"{name}.mp3")
    full_message = f"{name}{massage}"
    logger.info(f"正在为 {name} 生成语音: {voice_file}")
    while True:
        success = make_tts_kedaxunfei(
            logger,
            GlobalConf.app_id,
            GlobalConf.api_key,
            GlobalConf.vcn,
            full_message,
            voice_file
        )
        if success:
            break
    if os.path.exists(voice_file):
        audio = AudioSegment.from_mp3(voice_file)
        wav_path = os.path.join(temp_mp3_dir, f"{name}.wav")
        audio.export(wav_path, format="wav")
        logger.info(f"成功生成WAV文件: {wav_path}")
        # 删除mp3文件，只留wav文件
        os.remove(voice_file)
    else:
        logger.error(f"生成MP3文件失败: {voice_file}")

def generate_audio(name):
    try:
        excel_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "conf", "mess.xlsx")
        if not os.path.exists(excel_path):
            logger.error(f"Excel文件不存在: {excel_path}")
            return False

        workbook = load_workbook(excel_path)
        sheet = workbook.active
        messages = []
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if row and row[0]:
                messages.append(row[0])

        if not messages:
            logger.error("Excel文件中没有有效数据")
            return False

        random_message = random.choice(messages)
        down_voice(name, random_message)

        return True
    except Exception as e:
        logger.error(f"读取Excel文件失败: {e}")
        return False




# generate_audio("a7777")
