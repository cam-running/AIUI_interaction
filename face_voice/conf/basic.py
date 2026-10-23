# -*- coding: UTF-8 -*-
# @Author     zyb
# @FileName   basic.PY
# @Time       2025-05-29 23:21
import json
import inspect
import queue
import os
import logging
from logging.handlers import RotatingFileHandler
from aip import AipFace


class GlobalConf:
    identify_sta = "online"
    volume_num = 80
    action_num = 0
    action_status = False
    through_1 = 7
    through_2 = 7
    interval_time = 5
    action_code_1 = 2015
    action_code_2 = 2016
    interval_code = 0
    identifications = 3
    identifications_num = 5
    server_queue = queue.Queue()
    face_queue = queue.Queue()
    tmp_queue = queue.Queue()
    identify_queue = queue.Queue()
    action_queue = queue.Queue()
    excel_queue = queue.Queue()
    show_queue = queue.Queue()
    attendance_queue = queue.Queue()
    exist_data = {}
    all_users = []
    app_id = "352cdc21"
    api_key = "45b5f6a1a4374409e6d72107e365ec1e"
    vcn = "x4_lingfeizhe_oral"
    #face_host = "192.168.123.56"
    face_host = "192.168.123.38"
    robot_host = "192.168.1.220"
    robot_action = "AUTO_DROP_HANDSHAKE"
    robot_params = ["1.5", "5", "1.5"]
    action_url = "http://192.168.123.164:8839/dbqa_app/pub_webrtc"
    attendance_url = "http://192.168.123.164:8839/rest_api/sign_in_callback"
    home_path = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dest_path = os.path.join(home_path, "tmp")
    result_path = os.path.join(home_path, "result")
    tmp_path = os.path.join(home_path, "conf", "temp_mp3")
    voice_path = os.path.join(home_path, "voice")
    log_path = os.path.join(home_path, "logs")
    known_face_encodings = None
    known_face_names = None
    excel_name = "data_group7.xlsx"
    excel_file = os.path.join("name_file", excel_name)
    back_file = None
    sign_in_file = os.path.join(home_path, "sign_in.txt")
    min_distance = 0.35
    sleep_time = 5
    sink_status = False
    identify_status = False
    baidu_app_id = "117848461"
    baidu_api_key = "b8C3YBjkGyIbW8zK33yXwYkL"
    baidu_secret_key = "Q4MLzgfbv5zq7zI0rIBkrC8nMEAzBaSw"
    baidu_groups = "group7"
    baidu_gw_url = "http://www.zs2016.cloud:35001/add_users"
    baidu_score = 80
    show_face_cv = True
    clear_tmp = True
    group_list = ["group0", "group1", "group2", "group3", "group4", "group5", "group6", "group7", "group8", "group9"]

def modify_e(filename, clas_ve, values, keys):
    clas_ve = values
    filename.clas_ve = values
    file_path = inspect.getfile(filename)
    with open(file_path, "r", encoding='utf-8') as f:
        lines = f.readlines()
    for i, line in enumerate(lines):
        if keys in line and "=" in line:
            lines[i] = f'    {keys} = "{values}"\n'
            break
    with open(file_path, "w", encoding='utf-8') as f:
        f.writelines(lines)

def get_logger(log_name, log_dir="logs", level="INFO", max_bytes=7 * 1024 * 1024, backup_count=5):
    try:
        if not os.path.exists(log_dir):
            os.makedirs(log_dir)
    except OSError as e:
        log_dir = "."

    log_file = os.path.join(log_dir, f"{log_name}.log")
    logger = logging.getLogger(log_name)
    level_mapping = {
        "DEBUG": logging.DEBUG,
        "INFO": logging.INFO,
        "WARNING": logging.WARNING,
        "ERROR": logging.ERROR,
        "CRITICAL": logging.CRITICAL
    }
    logger.setLevel(level_mapping.get(level, logging.INFO))
    logger.propagate = False

    formatter = logging.Formatter(
        '%(asctime)s - %(name)s - %(levelname)s - %(message)s - Line: %(lineno)d'
    )
    if not logger.handlers:
        file_handler = RotatingFileHandler(
            log_file,
            maxBytes=max_bytes,
            backupCount=backup_count,
            encoding='utf-8'
        )
        file_handler.setFormatter(formatter)
        logger.addHandler(file_handler)
        console_handler = logging.StreamHandler()
        console_handler.setFormatter(formatter)
        logger.addHandler(console_handler)

    return logger

def read_config(group):
    CONFIG_FILE = os.path.join("conf", "group_name", f"{group}.json")
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return {}

def write_config(group_id, data):
    CONFIG_FILE = os.path.join("conf", "group_name", f"{group_id}.json")
    config = read_config(group_id)
    config[group_id] = data
    os.makedirs(os.path.dirname(CONFIG_FILE), exist_ok=True)
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=4)


def delete_user(group: str, name: str, userid: str):
    try:
        config_file_path = os.path.join("conf", "group_name", f"{group}.json")
        with open(config_file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if group in data:
            group_list = data[group]
            new_group_list = [user for user in group_list if not (user["name"] == name and user["userid"] == userid)]
            if len(new_group_list) == len(group_list):
                return {"message": "No matching user found", "success": False}
            data[group] = new_group_list
            with open(config_file_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=4)
            return {"message": f"User {name} successfully deleted", "success": True, "code": 200}
        else:
            return {"message": f"{group} group list not found", "success": False}
    except FileNotFoundError:
        return {"message": "The configuration file is not found", "success": False}
    except Exception as e:
        return {"message": f"An error occurred: {str(e)}", "success": False}


def add_u_json(group: str, name: str, userid: str):
    try:
        config_file_path = os.path.join("conf", "group_name", f"{group}.json")
        os.makedirs(os.path.dirname(config_file_path), exist_ok=True)
        if os.path.exists(config_file_path):
            with open(config_file_path, "r", encoding="utf-8") as f:
                data = json.load(f)
        else:
            data = {group: []}
        data[group].append({"name": name, "userid": userid})
        with open(config_file_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=4)
        return {"message": "Add successfully", "success": True}

    except FileNotFoundError:
        return {"message": "The configuration file is not found", "success": False}
    except Exception as e:
        return {"message": f"An error occurred: {str(e)}", "success": False}


def get_all_user(group: str):
    try:
        config_file_path = os.path.join("conf", "group_name", f"{group}.json")
        with open(config_file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        names = [user["name"] for user in data.get(group, [])]
        return names
    except FileNotFoundError:
        return {"code": 404, "message": f"{group}.json file not found", "data": []}
    except json.JSONDecodeError:
        return {"code": 500, "message": "The configuration file format is incorrect", "data": []}
    except Exception as e:
        return {"code": 500, "message": f"Failed to get the user list: {str(e)}", "data": []}

def count_first_column():
    try:
        from openpyxl import load_workbook
        excel_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "name_file", GlobalConf.excel_name)
        wb = load_workbook(excel_path)
        sheet = wb.active
        count = 0
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if row[0]:
                count += 1
        return count
    except Exception as e:
        logging.error(f"Error in the first column of statistics Excel: {str(e)}")
        return 0


def check_name_exists(name, group):
    try:
        from openpyxl import load_workbook, Workbook
        excel_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                 "name_file", f"data_{group}.xlsx")
        if not os.path.exists(excel_path):
            os.makedirs(os.path.dirname(excel_path), exist_ok=True)
            wb = Workbook()
            wb.save(excel_path)
            return False
        wb = load_workbook(excel_path)
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            if row and row[0] == name:
                return True
        return False
    except Exception as e:
        logging.error(f"Error when checking if the name exists: {str(e)}")
        return False

def get_user_time(name):
    conf_dir = os.path.dirname(__file__)
    parent_dir = os.path.dirname(conf_dir)
    file_path = os.path.join(parent_dir, "sign_in.txt")
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
            data = json.loads(content)
            user_data = data.get(name, {})
            time_value = user_data.get("time", 0)
            num_value = user_data.get("num", 0)
            if num_value == GlobalConf.identifications_num:
                return time_value, False
            return time_value, True
    except Exception as e:
        return 0, False

if __name__ == "__main__":
    pass
